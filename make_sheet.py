#!/usr/bin/env python3
"""Gera a planilha de de-para (+ aba Log) a partir de plan.json e log.jsonl.

Uso: make_sheet.py <pasta de trabalho> <saída .xlsx>
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

LINK = "https://drive.google.com/file/d/{}/view"


def main():
    work, out = Path(sys.argv[1]), Path(sys.argv[2])
    plan = json.loads((work / "plan.json").read_text())
    log = [json.loads(l) for l in (work / "log.jsonl").read_text().splitlines()]
    ids = {e["arquivo"]: e["id"] for e in log if e.get("status") == "ok"}

    wb = Workbook()
    ws = wb.active
    ws.title = "De-para"
    ws.append(["Arquivo original", "Novo nome", "Pasta", "Duração (s)", "Motivo", "Transcrição"])
    for name, p in sorted(plan.items(), key=lambda kv: (kv[1]["pasta"], kv[0])):
        ws.append([name, p.get("novo_nome", name), p["pasta"], p.get("duracao"),
                   p.get("motivo", ""), p.get("transcricao", "")])
        cell = ws.cell(ws.max_row, 2)
        if name in ids:
            cell.hyperlink = LINK.format(ids[name])
            cell.font = Font(color="1155CC", underline="single")
    for col, width in zip("ABCDEF", (22, 34, 30, 10, 60, 110)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    lg = wb.create_sheet("Log")
    lg.append(["Quando", "Arquivo", "Pasta", "Status", "Erro"])
    for e in log:
        lg.append([e["quando"], e["arquivo"], e.get("pasta"), e["status"], e.get("erro", "")])
    for col, width in zip("ABCDE", (20, 22, 34, 8, 60)):
        lg.column_dimensions[col].width = width

    for sheet in (ws, lg):
        for c in sheet[1]:
            c.font = Font(bold=True)
    wb.save(out)


if __name__ == "__main__":
    main()
